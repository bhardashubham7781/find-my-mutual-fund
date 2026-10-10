#!/usr/bin/env python3
"""Audit how observed equity industries map to the product's standard sectors.

Produces a review CSV that includes all configured labels, observed counts, and
unmapped/blank industry labels. This is an audit report only; it does not change
scheme eligibility or publish data.
"""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import process_portfolios as processor

PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_RAW_DIR = PROJECT_DIR / "data" / "raw"
DEFAULT_MAPPING = PROJECT_DIR / "config" / "sector_mapping.csv"
DEFAULT_REPORT = PROJECT_DIR / "reports" / "sector_mapping_audit.csv"
DEFAULT_SUMMARY = PROJECT_DIR / "reports" / "sector_mapping_audit_summary.json"


def collect_workbooks(input_dir: Path) -> list[Path]:
    return sorted(
        p for p in input_dir.rglob("*")
        if p.is_file()
        and p.suffix.lower() in {".xlsx", ".xlsm"}
        and not p.name.startswith("~$")
    )


def add_usage(usage: dict[str, dict[str, Any]], label: str, scheme: str, mapped_sector: str | None) -> None:
    raw = label.strip() if label and label.strip() else "<blank industry>"
    key = processor.normalise(raw)
    item = usage.setdefault(key, {
        "source_industry": raw,
        "occurrence_count": 0,
        "example_schemes": set(),
        "files": set(),
        "mapped_sector": mapped_sector or "",
    })
    item["occurrence_count"] += 1
    if scheme:
        item["example_schemes"].add(scheme)
    if mapped_sector is not None:
        item["mapped_sector"] = mapped_sector


def audit_sbi_workbook(workbook: Any, mapping: dict[str, str], usage: dict[str, dict[str, Any]], source_file: str) -> tuple[int, list[str]]:
    warnings: list[str] = []
    index_sheet = processor.find_index_sheet(workbook)
    if index_sheet is None:
        return 0, [f"{source_file}: no Index sheet"]
    entries = processor.build_index_rows(index_sheet)
    if not entries:
        return 0, [f"{source_file}: scheme index could not be read"]
    processed = 0
    for entry in entries:
        sheet_name = processor.find_sheet_name(workbook, entry["short_code"])
        if not sheet_name:
            warnings.append(f"{source_file}: missing scheme tab {entry['short_code']}")
            continue
        sheet = workbook[sheet_name]
        header = processor.find_main_table_header(sheet)
        if not header:
            warnings.append(f"{source_file}: no holdings header on {sheet_name}")
            continue
        header_row, instrument_col, percent_col, industry_col, market_col = header
        scheme_name = entry.get("scheme_name") or sheet_name
        in_equity = False
        for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, values_only=True):
            label = processor.text_value(row[instrument_col]) if instrument_col < len(row) else ""
            if not label:
                continue
            key = processor.normalise(label)
            if key.startswith("grand total") and "aum" in key:
                break
            if key == "equity and equity related" or key.startswith("equity and equity related"):
                in_equity = True
                continue
            if key == "debt instruments" or key.startswith("debt instruments"):
                in_equity = False
                continue
            if key == "total" or key.startswith("total ") or "derivatives total" in key:
                continue
            pct = processor.parse_number(row[percent_col]) if percent_col < len(row) else None
            if pct is None:
                continue
            if market_col is not None and market_col < len(row):
                market = processor.parse_number(row[market_col])
                if market is None and processor.text_value(row[market_col]).upper() in {"NIL", "N.A.", "N/A"}:
                    continue
            if not in_equity:
                continue
            raw_industry = processor.text_value(row[industry_col]) if industry_col is not None and industry_col < len(row) else ""
            mapped = mapping.get(processor.normalise(raw_industry)) if raw_industry else None
            add_usage(usage, raw_industry, scheme_name, mapped)
        processed += 1
    return processed, warnings


def audit_hdfc_workbook(workbook: Any, mapping: dict[str, str], usage: dict[str, dict[str, Any]], source_file: str) -> tuple[int, list[str]]:
    sheet = next((s for s in workbook.worksheets if processor.normalise(s.title) == "hdfceq"), None)
    if sheet is None:
        return 0, [f"{source_file}: no HDFCEQ sheet"]
    header = processor.find_hdfc_main_table_header(sheet)
    if not header:
        return 0, [f"{source_file}: HDFC holdings header not recognised"]
    header_row, instrument_col, industry_col, percent_col = header
    scheme_name, _ = processor.extract_hdfc_scheme_header(sheet)
    scheme_name = scheme_name or source_file
    in_equity = False
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, values_only=True):
        section_label = processor.normalise(row[1]) if len(row) > 1 else ""
        if section_label == "grand total":
            break
        if section_label == "equity and equity related":
            in_equity = True
            continue
        if section_label in {"debt instruments", "money market instruments", "others"}:
            in_equity = False
            continue
        if section_label in {"sub total", "total"}:
            continue
        instrument = processor.text_value(row[instrument_col]) if instrument_col < len(row) else ""
        if not instrument:
            continue
        pct = processor.parse_number(row[percent_col]) if percent_col < len(row) else None
        if pct is None or not in_equity:
            continue
        raw_industry = processor.text_value(row[industry_col]) if industry_col < len(row) else ""
        mapped = mapping.get(processor.normalise(raw_industry)) if raw_industry else None
        add_usage(usage, raw_industry, scheme_name, mapped)
    return 1, []


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a review-only audit of industry-to-sector mappings used in raw AMC workbooks.")
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_RAW_DIR)
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING)
    parser.add_argument("--output", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY)
    args = parser.parse_args()

    mapping = processor.read_mapping(args.mapping)
    raw_rows: list[dict[str, str]] = []
    with args.mapping.open("r", encoding="utf-8-sig", newline="") as handle:
        raw_rows = [row for row in csv.DictReader(handle) if row.get("source_industry", "").strip()]
    usage: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []
    workbook_count = 0
    source_schemes = 0

    paths = collect_workbooks(args.input_dir)
    if not paths:
        print(f"ERROR: no .xlsx/.xlsm files in {args.input_dir}")
        return 2

    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if processor.find_index_sheet(workbook) is not None:
                processed, current_warnings = audit_sbi_workbook(workbook, mapping, usage, path.name)
            elif any(processor.normalise(s.title) == "hdfceq" for s in workbook.worksheets):
                processed, current_warnings = audit_hdfc_workbook(workbook, mapping, usage, path.name)
            else:
                processed, current_warnings = 0, [f"{path.name}: unsupported layout for sector mapping audit"]
            workbook_count += 1 if processed else 0
            source_schemes += processed
            warnings.extend(current_warnings)
        finally:
            workbook.close()

    configured = {processor.normalise(row["source_industry"]): row for row in raw_rows}
    keys = set(configured) | set(usage)
    output_rows: list[dict[str, Any]] = []
    totals = Counter()
    for key in sorted(keys):
        cfg = configured.get(key, {})
        info = usage.get(key, {})
        source_industry = (info.get("source_industry") or cfg.get("source_industry") or "<blank industry>")
        sector = mapping.get(key, "")
        count = int(info.get("occurrence_count", 0))
        if key == processor.normalise("<blank industry>"):
            status = "REVIEW_BLANK_INDUSTRY"
        elif key not in mapping:
            status = "UNMAPPED_DEFAULTS_TO_OTHERS"
        elif sector == "Others":
            status = "REVIEW_MAPPING_TO_OTHERS"
        elif count == 0:
            status = "CONFIGURED_NOT_OBSERVED_IN_CURRENT_FILES"
        else:
            status = "MAPPED_OBSERVED_REVIEW_SAMPLE"
        totals[status] += 1
        output_rows.append({
            "source_industry": source_industry,
            "standard_sector": sector or "Others (unmapped fallback)",
            "occurrence_count": count,
            "example_schemes": " | ".join(sorted(info.get("example_schemes", set()))[:8]),
            "mapping_notes": cfg.get("notes", ""),
            "review_status": status,
        })

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "source_industry", "standard_sector", "occurrence_count", "example_schemes", "mapping_notes", "review_status",
        ], lineterminator="\n")
        writer.writeheader()
        writer.writerows(output_rows)

    summary = {
        "generated_date": __import__("datetime").date.today().isoformat(),
        "source_workbooks_found": len(paths),
        "workbooks_recognised": workbook_count,
        "scheme_tabs_or_workbooks_scanned": source_schemes,
        "configured_mapping_labels": len(configured),
        "distinct_observed_industry_labels": len(usage),
        "review_status_counts": dict(totals),
        "warnings": warnings,
        "important": "This report helps review the sector mapping. A mapped label is not automatically correct; review each assignment, especially all labels currently mapped to Others, before trusting or publishing real results.",
    }
    args.summary.parent.mkdir(parents=True, exist_ok=True)
    args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"Mapping audit CSV: {args.output}")
    print(f"Summary JSON: {args.summary}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
