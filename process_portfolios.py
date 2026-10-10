#!/usr/bin/env python3
"""Build sector-level mutual fund data from supported AMC portfolio workbooks.

Initial adapter: SBI Mutual Fund "All Schemes Monthly Portfolio" .xlsx layout.
The script is intentionally conservative: unknown equity industry labels are
assigned to Others so totals remain visible, and separately written to an
unmapped-industries report for review.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import tempfile
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

STANDARD_SECTORS = [
    "IT",
    "Banking",
    "Financial Services",
    "Pharma & Healthcare",
    "Auto",
    "FMCG",
    "Energy/Oil & Gas",
    "Power",
    "Metals & Mining",
    "Capital Goods",
    "Construction/Infra",
    "Telecom",
    "Consumer Durables",
    "Realty",
    "Chemicals",
    "Others",
]

PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_RAW_DIR = PROJECT_DIR / "data" / "raw"
DEFAULT_MAPPING = PROJECT_DIR / "config" / "sector_mapping.csv"
DEFAULT_SCHEME_MASTER = PROJECT_DIR / "config" / "scheme_master.csv"
DEFAULT_DATA_OUT = PROJECT_DIR / "data" / "processed"
DEFAULT_REPORTS_OUT = PROJECT_DIR / "reports"


def text_value(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalise(value: Any) -> str:
    """Case- and punctuation-insensitive key used for industry/scheme matching."""
    text = text_value(value).lower().replace("&", " and ")
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"\^+", "", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def canonical_scheme_name(value: Any) -> str:
    """Normalise common scheme-name variants and remove plan/option suffixes."""
    raw_text = text_value(value)
    # Some AMFI names retain a legacy scheme name in parentheses after a merger/rename.
    raw_text = re.sub(r"\s*\(erstwhile\s+[^)]*\)", "", raw_text, flags=re.IGNORECASE)
    text = normalise(raw_text)
    # AMCs and AMFI sometimes join the same category words differently:
    # e.g. "Smallcap" in a portfolio workbook vs "Small Cap" in the AMFI feed.
    for joined, spaced in [
        ("smallcap", "small cap"),
        ("midcap", "mid cap"),
        ("largecap", "large cap"),
        ("flexicap", "flexi cap"),
        ("multicap", "multi cap"),
    ]:
        text = re.sub(r"\b" + joined + r"\b", spaced, text)
    # Remove plan labels wherever they occur; then remove common trailing options.
    text = re.sub(r"\b(direct plan|regular plan|institutional plan)\b", " ", text)
    suffixes = [
        "growth option", "growth", "idcw payout", "idcw reinvestment", "idcw",
        "dividend payout", "dividend reinvestment", "dividend option", "dividend",
        "payout", "reinvestment", "bonus option", "bonus",
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if text == suffix:
                text = ""
                changed = True
            elif text.endswith(" " + suffix):
                text = text[: -(len(suffix) + 1)].strip()
                changed = True
    return re.sub(r"\s+", " ", text).strip()


def normalise_amc(value: Any) -> str:
    text = normalise(value)
    text = re.sub(r"\b(mutual fund|asset management company|asset management)\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_number(value: Any) -> float | None:
    """Parse a number. SBI uses # for values below 0.005%; treat that as zero."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = text_value(value)
    if not text:
        return None
    if text == "#":
        return 0.0
    if text.upper() in {"NIL", "N.A.", "NA", "N/A", "-", "—"}:
        return None
    text = text.replace(",", "").replace("%", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def parse_date(value: Any, workbook_epoch: Any = None) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and workbook_epoch is not None:
        try:
            converted = from_excel(value, workbook_epoch)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except Exception:
            pass
    text = text_value(value)
    if not text:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %B %Y", "%b %d,%Y", "%b %d, %Y", "%B %d,%Y", "%B %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def bool_value(value: Any) -> bool:
    return text_value(value).lower() in {"true", "yes", "y", "1", "include", "included"}


def read_mapping(path: Path) -> dict[str, str]:
    if not path.exists():
        raise FileNotFoundError(f"Sector mapping file not found: {path}")
    mapping: dict[str, str] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"source_industry", "standard_sector"}
        if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
            raise ValueError("sector_mapping.csv must have columns: source_industry, standard_sector, notes")
        for row in reader:
            raw = text_value(row.get("source_industry"))
            sector = text_value(row.get("standard_sector"))
            if not raw:
                continue
            if sector not in STANDARD_SECTORS:
                raise ValueError(f"Invalid standard sector '{sector}' for source industry '{raw}'. Allowed: {', '.join(STANDARD_SECTORS)}")
            mapping[normalise(raw)] = sector
    return mapping


def read_scheme_master(path: Path) -> list[dict[str, str]]:
    """Load the editable crosswalk. An empty template is allowed during setup."""
    if not path.exists():
        return []
    records: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"scheme_name", "category", "plan", "option", "include_v1"}
        if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
            raise ValueError("scheme_master.csv must have columns: amc, scheme_name, amfi_scheme_code, category, plan, option, include_v1")
        for row in reader:
            if not text_value(row.get("scheme_name")):
                continue
            record = {key: text_value(value) for key, value in row.items() if key}
            record["_canonical_name"] = canonical_scheme_name(record.get("scheme_name"))
            record["_amc_key"] = normalise_amc(record.get("amc"))
            record["_include_v1"] = "true" if bool_value(record.get("include_v1")) else "false"
            records.append(record)
    return records


def find_index_sheet(workbook: Any) -> Any | None:
    for sheet in workbook.worksheets:
        if normalise(sheet.title) == "index":
            return sheet
    return None


def build_index_rows(index_sheet: Any) -> list[dict[str, str]]:
    rows = list(index_sheet.iter_rows(values_only=True))
    header_at = None
    code_col = short_col = name_col = None
    for i, row in enumerate(rows[:30]):
        normalized = [normalise(v) for v in row]
        for j, heading in enumerate(normalized):
            if heading == "scheme code":
                code_col = j
            elif heading == "scheme short code":
                short_col = j
            elif heading == "scheme name":
                name_col = j
        if code_col is not None and short_col is not None and name_col is not None:
            header_at = i
            break
    if header_at is None:
        return []

    result = []
    for row in rows[header_at + 1 :]:
        if len(row) <= max(code_col, short_col, name_col):
            continue
        short_code = text_value(row[short_col])
        scheme_name = text_value(row[name_col])
        code_value = row[code_col]
        if not short_code or not scheme_name:
            continue
        if isinstance(code_value, (int, float)) and float(code_value).is_integer():
            amc_code = f"{int(code_value):03d}"
        else:
            amc_code = text_value(code_value)
        result.append({"short_code": short_code, "amc_scheme_code": amc_code, "scheme_name": scheme_name})
    return result


def find_sheet_name(workbook: Any, preferred: str) -> str | None:
    if preferred in workbook.sheetnames:
        return preferred
    by_normalised = {normalise(name): name for name in workbook.sheetnames}
    return by_normalised.get(normalise(preferred))


def get_scheme_header(sheet: Any) -> tuple[str, date | None, str]:
    scheme_name = ""
    amc = ""
    as_of_date: date | None = None
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True):
        for i, value in enumerate(row):
            text = text_value(value)
            key = normalise(text)
            if "scheme name" == key:
                for next_value in row[i + 1 :]:
                    if text_value(next_value):
                        scheme_name = text_value(next_value)
                        break
            if "portfolio statement as on" in key or key.endswith("statement as on"):
                for next_value in row[i + 1 :]:
                    parsed = parse_date(next_value, sheet.parent.epoch)
                    if parsed:
                        as_of_date = parsed
                        break
            if not amc and ("mutual fund" in key or "asset management" in key):
                amc = text
    return scheme_name, as_of_date, amc


def find_main_table_header(sheet: Any) -> tuple[int, int, int, int | None, int | None] | None:
    """Return (header row, instrument col, percent col, industry col, market value col), 1-based row and 0-based columns."""
    for r, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 40), values_only=True), start=1):
        instrument_col = percent_col = industry_col = market_col = None
        for c, value in enumerate(row):
            text = re.sub(r"\s+", " ", text_value(value)).strip().lower()
            squashed = re.sub(r"\s+", "", text)
            if "name of the instrument" in text:
                instrument_col = c
            if "% to aum" in text or "%toaum" in squashed:
                percent_col = c
            if "industry" in text:
                industry_col = c
            if "market value" in text:
                market_col = c
        if instrument_col is not None and percent_col is not None:
            return r, instrument_col, percent_col, industry_col, market_col
    return None


def process_scheme_sheet(
    workbook: Any,
    sheet: Any,
    index_entry: dict[str, str],
    mapping: dict[str, str],
    source_file: str,
    unmapped: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    sheet_scheme_name, as_of, amc_name = get_scheme_header(sheet)
    scheme_name = sheet_scheme_name or index_entry["scheme_name"]
    if not amc_name:
        amc_name = "Unknown AMC"
    header = find_main_table_header(sheet)
    sectors = {sector: 0.0 for sector in STANDARD_SECTORS}
    holdings_count = 0
    unmapped_for_scheme: set[str] = set()
    reported_total_pct = None
    in_equity_section = False

    if header is None:
        flags = ["MAIN_HOLDINGS_HEADER_NOT_FOUND"]
        return {
            "amc": amc_name, "scheme_name": scheme_name, "amc_scheme_code": index_entry["amc_scheme_code"],
            "source_file": source_file, "source_sheet": sheet.title,
            "as_of_date": as_of.isoformat() if as_of else None,
            "as_of_month": as_of.strftime("%Y-%m") if as_of else None,
            "sector_exposure": sectors, "sector_total_pct": 0.0, "reported_grand_total_pct": None,
            "holdings_count": 0, "unmapped_industries": [], "quality_flags": flags,
        }

    header_row, instrument_col, percent_col, industry_col, market_col = header
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, values_only=True):
        instrument = row[instrument_col] if instrument_col < len(row) else None
        label = text_value(instrument)
        key = normalise(label)

        if not label:
            continue
        if key.startswith("grand total") and "aum" in key:
            if percent_col < len(row):
                reported_total_pct = parse_number(row[percent_col])
            break
        if key == "equity and equity related" or key.startswith("equity and equity related"):
            in_equity_section = True
            continue
        if key == "debt instruments" or key.startswith("debt instruments"):
            in_equity_section = False
            continue

        if key == "total" or key.startswith("total ") or "derivatives total" in key:
            continue
        if "name of the instrument" in key:
            continue

        pct = parse_number(row[percent_col]) if percent_col < len(row) else None
        if pct is None:
            continue
        # A holding/instrument row must carry a numeric or SBI '#' percentage.
        # The instrument name must be present; section captions and NIL rows are skipped.
        if market_col is not None and market_col < len(row):
            market_value = parse_number(row[market_col])
            if market_value is None and text_value(row[market_col]).upper() in {"NIL", "N.A.", "N/A"}:
                continue
        if not label:
            continue

        holdings_count += 1
        if not in_equity_section:
            sector = "Others"
        else:
            raw_industry = text_value(row[industry_col]) if industry_col is not None and industry_col < len(row) else ""
            mapped_sector = mapping.get(normalise(raw_industry)) if raw_industry else None
            if mapped_sector:
                sector = mapped_sector
            else:
                sector = "Others"
                raw_key = raw_industry if raw_industry else "<blank industry>"
                unmapped_for_scheme.add(raw_key)
                item = unmapped.setdefault(raw_key, {"occurrences": 0, "schemes": set()})
                item["occurrences"] += 1
                item["schemes"].add(scheme_name)
        sectors[sector] += pct

    sector_total = sum(sectors.values())
    flags: list[str] = []
    if holdings_count == 0:
        flags.append("MISSING_HOLDINGS_DATA")
    if as_of is None:
        flags.append("MISSING_AS_OF_DATE")
    if sector_total < 95 or sector_total > 101:
        flags.append("SECTOR_TOTAL_OUTSIDE_95_101")
    if unmapped_for_scheme:
        flags.append("UNMAPPED_EQUITY_INDUSTRY")

    return {
        "amc": amc_name,
        "scheme_name": scheme_name,
        "amc_scheme_code": index_entry["amc_scheme_code"],
        "source_file": source_file,
        "source_sheet": sheet.title,
        "as_of_date": as_of.isoformat() if as_of else None,
        "as_of_month": as_of.strftime("%Y-%m") if as_of else None,
        "sector_exposure": {sector: (round(value, 4) if abs(value) >= 0.00005 else 0.0) for sector, value in sectors.items()},
        "sector_total_pct": round(sector_total, 4),
        "reported_grand_total_pct": round(reported_total_pct, 4) if reported_total_pct is not None else None,
        "holdings_count": holdings_count,
        "unmapped_industries": sorted(unmapped_for_scheme),
        "quality_flags": flags,
    }


def find_hdfc_main_table_header(sheet: Any) -> tuple[int, int, int, int] | None:
    """Find HDFC's individual-scheme holdings header: instrument, industry and % to NAV."""
    for r, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 40), values_only=True), start=1):
        instrument_col = industry_col = percent_col = None
        for c, value in enumerate(row):
            raw = text_value(value).lower()
            key = normalise(value)
            if "name of the instrument" in key:
                instrument_col = c
            if "industry" in key:
                industry_col = c
            if "% to nav" in raw or key == "to nav":
                percent_col = c
        if instrument_col is not None and industry_col is not None and percent_col is not None:
            return r, instrument_col, industry_col, percent_col
    return None


def extract_hdfc_scheme_header(sheet: Any) -> tuple[str, date | None]:
    """Read the scheme name and 'Portfolio as on ...' date from an HDFC portfolio sheet."""
    title = text_value(sheet.cell(1, 1).value)
    # HDFC puts the scheme name followed by a parenthetical scheme description in row 1.
    scheme_name = re.sub(r"\s*\(.*$", "", title).strip()
    as_of: date | None = None
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True):
        for value in row:
            text = text_value(value)
            match = re.search(r"\b\d{1,2}-[A-Za-z]{3}-\d{4}\b", text)
            if match:
                parsed = parse_date(match.group(0), sheet.parent.epoch)
                if parsed:
                    as_of = parsed
                    break
        if as_of:
            break
    return scheme_name, as_of


def process_hdfc_scheme_sheet(
    sheet: Any,
    mapping: dict[str, str],
    source_file: str,
    unmapped: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Process one HDFC monthly portfolio workbook's HDFCEQ sheet."""
    scheme_name, as_of = extract_hdfc_scheme_header(sheet)
    header = find_hdfc_main_table_header(sheet)
    sectors = {sector: 0.0 for sector in STANDARD_SECTORS}
    holdings_count = 0
    unmapped_for_scheme: set[str] = set()
    reported_total_pct: float | None = None
    in_equity_section = False

    if header is None:
        flags = ["MAIN_HOLDINGS_HEADER_NOT_FOUND"]
        return {
            "amc": "HDFC Mutual Fund", "scheme_name": scheme_name or source_file,
            "amc_scheme_code": "", "source_file": source_file, "source_sheet": sheet.title,
            "as_of_date": as_of.isoformat() if as_of else None,
            "as_of_month": as_of.strftime("%Y-%m") if as_of else None,
            "sector_exposure": sectors, "sector_total_pct": 0.0,
            "reported_grand_total_pct": None, "holdings_count": 0,
            "unmapped_industries": [], "quality_flags": flags,
        }

    header_row, instrument_col, industry_col, percent_col = header
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, values_only=True):
        # HDFC section captions and totals are generally in column B, unlike its
        # holding descriptions which are in the "Name Of the Instrument" column.
        section_label = normalise(row[1]) if len(row) > 1 else ""

        if section_label == "grand total":
            reported_total_pct = parse_number(row[percent_col]) if percent_col < len(row) else None
            break
        if section_label == "equity and equity related":
            in_equity_section = True
            continue
        if section_label in {"debt instruments", "money market instruments", "others"}:
            in_equity_section = False
            continue
        if section_label in {"sub total", "total"}:
            continue

        instrument = row[instrument_col] if instrument_col < len(row) else None
        label = text_value(instrument)
        if not label:
            continue
        pct = parse_number(row[percent_col]) if percent_col < len(row) else None
        if pct is None:
            continue

        holdings_count += 1
        if not in_equity_section:
            # Bonds, money-market instruments and net current assets are counted in
            # Others so the sector breakdown reconciles to the portfolio's % of NAV.
            sectors["Others"] += pct
            continue

        raw_industry = text_value(row[industry_col]) if industry_col < len(row) else ""
        mapped_sector = mapping.get(normalise(raw_industry)) if raw_industry else None
        if mapped_sector:
            sector = mapped_sector
        else:
            sector = "Others"
            raw_key = raw_industry if raw_industry else "<blank industry>"
            unmapped_for_scheme.add(raw_key)
            item = unmapped.setdefault(raw_key, {"occurrences": 0, "schemes": set()})
            item["occurrences"] += 1
            item["schemes"].add(scheme_name or source_file)
        sectors[sector] += pct

    sector_total = sum(sectors.values())
    flags: list[str] = []
    if holdings_count == 0:
        flags.append("MISSING_HOLDINGS_DATA")
    if as_of is None:
        flags.append("MISSING_AS_OF_DATE")
    if sector_total < 95 or sector_total > 101:
        flags.append("SECTOR_TOTAL_OUTSIDE_95_101")
    if unmapped_for_scheme:
        flags.append("UNMAPPED_EQUITY_INDUSTRY")

    return {
        "amc": "HDFC Mutual Fund",
        "scheme_name": scheme_name or source_file,
        "amc_scheme_code": "",
        "source_file": source_file,
        "source_sheet": sheet.title,
        "as_of_date": as_of.isoformat() if as_of else None,
        "as_of_month": as_of.strftime("%Y-%m") if as_of else None,
        "sector_exposure": {sector: (round(value, 4) if abs(value) >= 0.00005 else 0.0) for sector, value in sectors.items()},
        "sector_total_pct": round(sector_total, 4),
        "reported_grand_total_pct": round(reported_total_pct, 4) if reported_total_pct is not None else None,
        "holdings_count": holdings_count,
        "unmapped_industries": sorted(unmapped_for_scheme),
        "quality_flags": flags,
    }


def find_icici_main_table_header(sheet: Any) -> tuple[int, int, int, int] | None:
    """Return header row, instrument column, industry/rating column, and % to NAV column."""
    for r, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True), start=1):
        instrument_col = industry_col = percent_col = None
        for c, value in enumerate(row):
            raw = text_value(value).lower().replace(" ", "")
            key = normalise(value)
            if "instrument name" in key and ("company" in key or "issuer" in key):
                instrument_col = c
            if "industry rating" in key:
                industry_col = c
            if "%tonav" in raw:
                percent_col = c
        if instrument_col is not None and industry_col is not None and percent_col is not None:
            return r, instrument_col, industry_col, percent_col
    return None


def extract_icici_scheme_header(sheet: Any) -> tuple[str, date | None]:
    """Extract scheme title and disclosure date from ICICI's individual-scheme workbook."""
    scheme_name = ""
    as_of: date | None = None
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True):
        for value in row:
            raw = text_value(value)
            key = normalise(raw)
            if not scheme_name and key.startswith("icici prudential") and "mutual fund" not in key and "portfolio as on" not in key:
                scheme_name = raw
            if as_of is None and "portfolio as on" in key:
                match = re.search(r"([A-Za-z]{3,9}\s+\d{1,2},\s*\d{4})", raw)
                if match:
                    as_of = parse_date(match.group(1))
        if scheme_name and as_of:
            break
    return scheme_name, as_of


def process_icici_scheme_sheet(
    sheet: Any,
    mapping: dict[str, str],
    source_file: str,
    unmapped: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Process ICICI's single-scheme portfolio. Its % to NAV values are fractions, not percentage points."""
    scheme_name, as_of = extract_icici_scheme_header(sheet)
    header = find_icici_main_table_header(sheet)
    sectors = {sector: 0.0 for sector in STANDARD_SECTORS}
    holdings_count = 0
    unmapped_for_scheme: set[str] = set()
    reported_total_pct: float | None = None
    in_equity_section = False

    if header is None:
        return {
            "amc": "ICICI Prudential Mutual Fund", "scheme_name": scheme_name or source_file,
            "amc_scheme_code": "", "source_file": source_file, "source_sheet": sheet.title,
            "as_of_date": as_of.isoformat() if as_of else None,
            "as_of_month": as_of.strftime("%Y-%m") if as_of else None,
            "sector_exposure": sectors, "sector_total_pct": 0.0,
            "reported_grand_total_pct": None, "holdings_count": 0,
            "unmapped_industries": [], "quality_flags": ["MAIN_HOLDINGS_HEADER_NOT_FOUND"],
        }

    header_row, instrument_col, industry_col, percent_col = header
    # ICICI's first sheet has the full portfolio. The separate Derivative sheet is
    # deliberately excluded; the main table stops at Total Net Assets before futures.
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, values_only=True):
        instrument = row[instrument_col] if instrument_col < len(row) else None
        label = text_value(instrument)
        key = normalise(label)
        if not label:
            continue

        raw_pct = parse_number(row[percent_col]) if percent_col < len(row) else None
        if key == "total net assets" or key.startswith("total net assets "):
            reported_total_pct = raw_pct * 100 if raw_pct is not None else None
            break
        if key.startswith("equity and equity related instruments"):
            in_equity_section = True
            continue
        if key in {"debt instruments", "money market instruments", "others"}:
            in_equity_section = False
            # These are sub-total rows; their underlying holdings may follow.
            continue
        if key in {
            "listed awaiting listing on stock exchanges", "unlisted",
            "securitized debt instruments", "term deposits",
            "deposits maturity not exceeding 91 days", "deposits placed as margin",
            "certificate of deposits", "commercial papers", "treasury bills",
            "units of real estate investment trust reits",
            "units of an alternative investment fund aif",
        }:
            continue
        if "details of stock future" in key or "details of index future" in key:
            break

        # Percentage cells are fractions of NAV (e.g. 0.089 means 8.9%).
        pct = raw_pct * 100 if raw_pct is not None else None
        if pct is None:
            continue

        # ICICI reports subtotal/group rows with a blank ISIN or instrument
        # identifier. Skip them to avoid counting a subtotal and its holdings.
        # TREPS, Cash Margin and Net Current Assets are actual balance-sheet
        # components and must be included even though their ISIN cells are blank.
        isin = text_value(row[2]) if len(row) > 2 else ""
        if not isin and key not in {"treps", "cash margin derivatives", "net current assets"}:
            continue

        if key == "net current assets":
            in_equity_section = False
        if key == "treps":
            in_equity_section = False
        if key == "cash margin derivatives":
            in_equity_section = False

        holdings_count += 1
        if not in_equity_section:
            sectors["Others"] += pct
            continue

        raw_industry = text_value(row[industry_col]) if industry_col < len(row) else ""
        mapped_sector = mapping.get(normalise(raw_industry)) if raw_industry else None
        if mapped_sector:
            sector = mapped_sector
        else:
            sector = "Others"
            raw_key = raw_industry if raw_industry else "<blank industry>"
            unmapped_for_scheme.add(raw_key)
            item = unmapped.setdefault(raw_key, {"occurrences": 0, "schemes": set()})
            item["occurrences"] += 1
            item["schemes"].add(scheme_name or source_file)
        sectors[sector] += pct

    sector_total = sum(sectors.values())
    flags: list[str] = []
    if holdings_count == 0:
        flags.append("MISSING_HOLDINGS_DATA")
    if as_of is None:
        flags.append("MISSING_AS_OF_DATE")
    if sector_total < 95 or sector_total > 101:
        flags.append("SECTOR_TOTAL_OUTSIDE_95_101")
    if unmapped_for_scheme:
        flags.append("UNMAPPED_EQUITY_INDUSTRY")

    return {
        "amc": "ICICI Prudential Mutual Fund",
        "scheme_name": scheme_name or source_file,
        "amc_scheme_code": "",
        "source_file": source_file,
        "source_sheet": sheet.title,
        "as_of_date": as_of.isoformat() if as_of else None,
        "as_of_month": as_of.strftime("%Y-%m") if as_of else None,
        "sector_exposure": {sector: (round(value, 4) if abs(value) >= 0.00005 else 0.0) for sector, value in sectors.items()},
        "sector_total_pct": round(sector_total, 4),
        "reported_grand_total_pct": round(reported_total_pct, 4) if reported_total_pct is not None else None,
        "holdings_count": holdings_count,
        "unmapped_industries": sorted(unmapped_for_scheme),
        "quality_flags": flags,
    }


def parse_supported_workbook(path: Path, mapping: dict[str, str], unmapped: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    warnings: list[str] = []
    index_sheet = find_index_sheet(workbook)
    if index_sheet is None:
        # HDFC publishes individual-scheme workbooks. Each workbook has a holdings
        # sheet commonly named HDFCEQ plus a separate derivatives disclosure sheet.
        hdfc_sheet = next(
            (sheet for sheet in workbook.worksheets if normalise(sheet.title) == "hdfceq"),
            None,
        )
        if hdfc_sheet is not None and find_hdfc_main_table_header(hdfc_sheet):
            record = process_hdfc_scheme_sheet(hdfc_sheet, mapping, path.name, unmapped)
            workbook.close()
            return [record], warnings

        # ICICI Prudential also publishes individual-scheme files. The main holdings
        # table has "Company/Issuer/Instrument Name", "Industry/Rating", and "% to Nav".
        # Exclude the separate derivatives sheet to avoid counting exposure twice.
        icici_sheet = next(
            (
                sheet for sheet in workbook.worksheets
                if normalise(sheet.title) not in {"derivative", "derivatives"}
                and find_icici_main_table_header(sheet)
            ),
            None,
        )
        if icici_sheet is not None:
            record = process_icici_scheme_sheet(icici_sheet, mapping, path.name, unmapped)
            workbook.close()
            return [record], warnings

        workbook.close()
        return [], [f"Unsupported workbook layout (no Index sheet or recognised HDFC/ICICI holdings sheet): {path.name}"]

    index_entries = build_index_rows(index_sheet)
    if not index_entries:
        workbook.close()
        return [], [f"Unsupported workbook layout (scheme index not recognised): {path.name}"]

    records: list[dict[str, Any]] = []
    for entry in index_entries:
        sheet_name = find_sheet_name(workbook, entry["short_code"])
        if not sheet_name:
            warnings.append(f"Missing scheme tab '{entry['short_code']}' in {path.name}")
            continue
        sheet = workbook[sheet_name]
        records.append(process_scheme_sheet(workbook, sheet, entry, mapping, path.name, unmapped))
    workbook.close()
    return records, warnings


def select_latest(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """If multiple months are present, retain the latest disclosure for each underlying scheme."""
    latest: dict[tuple[str, str], dict[str, Any]] = {}
    for record in records:
        key = (normalise_amc(record.get("amc")), normalise(record.get("scheme_name")))
        existing = latest.get(key)
        incoming_date = record.get("as_of_date") or "0000-00-00"
        existing_date = (existing or {}).get("as_of_date") or "0000-00-00"
        if existing is None or incoming_date >= existing_date:
            latest[key] = record
    return sorted(latest.values(), key=lambda r: (normalise_amc(r.get("amc")), normalise(r.get("scheme_name"))))


def master_candidates(record: dict[str, Any], master_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    base = canonical_scheme_name(record.get("scheme_name"))
    amc_key = normalise_amc(record.get("amc"))
    matching = [m for m in master_rows if m.get("_canonical_name") == base]
    same_amc = [m for m in matching if not m.get("_amc_key") or m.get("_amc_key") == amc_key]
    return same_amc or matching


def attach_scheme_master(records: list[dict[str, Any]], master_rows: list[dict[str, str]]) -> None:
    for record in records:
        candidates = master_candidates(record, master_rows)
        eligible = [m for m in candidates if m.get("_include_v1") == "true"]
        if eligible:
            selected = eligible[0]
            record["scheme_name"] = selected.get("scheme_name") or record["scheme_name"]
            record["amfi_scheme_code"] = selected.get("amfi_scheme_code") or None
            record["category"] = selected.get("category") or None
            record["plan"] = selected.get("plan") or None
            record["option"] = selected.get("option") or None
            record["include_v1"] = True
            record["scheme_master_status"] = "matched_eligible"
            if len(eligible) > 1:
                record["quality_flags"].append("AMBIGUOUS_SCHEME_MASTER_MATCH")
        elif candidates:
            # Metadata match exists, but the row has not been explicitly included in v1.
            selected = candidates[0]
            record["amfi_scheme_code"] = selected.get("amfi_scheme_code") or None
            record["category"] = selected.get("category") or None
            record["plan"] = selected.get("plan") or None
            record["option"] = selected.get("option") or None
            record["include_v1"] = False
            record["scheme_master_status"] = "matched_not_eligible"
        else:
            record["amfi_scheme_code"] = None
            record["category"] = None
            record["plan"] = None
            record["option"] = None
            record["include_v1"] = False
            record["scheme_master_status"] = "not_matched"
            record["quality_flags"].append("NO_SCHEME_MASTER_MATCH")


def add_missing_master_schemes(quality_rows: list[dict[str, Any]], records: list[dict[str, Any]], master_rows: list[dict[str, str]]) -> None:
    present_keys = {
        (normalise_amc(r.get("amc")), canonical_scheme_name(r.get("scheme_name")))
        for r in records
    }
    for master in master_rows:
        if master.get("_include_v1") != "true":
            continue
        key = (master.get("_amc_key", ""), master.get("_canonical_name", ""))
        # If an AMC key is supplied in the master, compare it; otherwise compare by scheme name.
        exists = any(
            present_name == key[1] and (not key[0] or present_amc == key[0])
            for present_amc, present_name in present_keys
        )
        if not exists:
            quality_rows.append({
                "amc": master.get("amc", ""),
                "scheme_name": master.get("scheme_name", ""),
                "amfi_scheme_code": master.get("amfi_scheme_code", ""),
                "category": master.get("category", ""),
                "plan": master.get("plan", ""),
                "option": master.get("option", ""),
                "source_file": "",
                "source_sheet": "",
                "as_of_date": "",
                "as_of_month": "",
                "sector_total_pct": "",
                "reported_grand_total_pct": "",
                "holdings_count": 0,
                "unmapped_industries": "",
                "scheme_master_status": "missing_portfolio_data",
                "include_v1": True,
                "quality_flags": "MISSING_PORTFOLIO_DATA",
            })


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def build_outputs(records: list[dict[str, Any]], master_rows: list[dict[str, str]], unmapped: dict[str, dict[str, Any]], warnings: list[str], data_out: Path, reports_out: Path) -> dict[str, Any]:
    records = select_latest(records)
    attach_scheme_master(records, master_rows)
    data_out.mkdir(parents=True, exist_ok=True)
    reports_out.mkdir(parents=True, exist_ok=True)

    funds_json = {
        "schema_version": 1,
        "disclaimer": "Data is from AMC monthly portfolio disclosures and may be delayed. This is not investment advice.",
        "sectors": STANDARD_SECTORS,
        "note": "Initial data build. Only records explicitly marked include_v1=true should be displayed in the version 1 screener.",
        "funds": records,
    }
    (data_out / "funds.json").write_text(json.dumps(funds_json, ensure_ascii=False, indent=2), encoding="utf-8")

    quality_rows: list[dict[str, Any]] = []
    for r in records:
        flags = list(r.get("quality_flags", []))
        quality_rows.append({
            "amc": r.get("amc", ""),
            "scheme_name": r.get("scheme_name", ""),
            "amc_scheme_code": r.get("amc_scheme_code", ""),
            "amfi_scheme_code": r.get("amfi_scheme_code") or "",
            "category": r.get("category") or "",
            "plan": r.get("plan") or "",
            "option": r.get("option") or "",
            "source_file": r.get("source_file", ""),
            "source_sheet": r.get("source_sheet", ""),
            "as_of_date": r.get("as_of_date") or "",
            "as_of_month": r.get("as_of_month") or "",
            "sector_total_pct": r.get("sector_total_pct", ""),
            "reported_grand_total_pct": r.get("reported_grand_total_pct") if r.get("reported_grand_total_pct") is not None else "",
            "holdings_count": r.get("holdings_count", 0),
            "unmapped_industries": " | ".join(r.get("unmapped_industries", [])),
            "scheme_master_status": r.get("scheme_master_status", ""),
            "include_v1": r.get("include_v1", False),
            "quality_flags": " | ".join(flags),
        })
    add_missing_master_schemes(quality_rows, records, master_rows)
    write_csv(reports_out / "data_quality.csv", quality_rows, [
        "amc", "scheme_name", "amc_scheme_code", "amfi_scheme_code", "category", "plan", "option",
        "source_file", "source_sheet", "as_of_date", "as_of_month", "sector_total_pct",
        "reported_grand_total_pct", "holdings_count", "unmapped_industries", "scheme_master_status",
        "include_v1", "quality_flags",
    ])

    unmapped_rows = []
    for industry, info in sorted(unmapped.items(), key=lambda item: item[0].lower()):
        unmapped_rows.append({
            "source_industry": industry,
            "occurrences": info["occurrences"],
            "example_schemes": " | ".join(sorted(info["schemes"])[:10]),
            "proposed_sector": "Others",
            "review_status": "REVIEW_MAPPING",
        })
    write_csv(reports_out / "unmapped_industries.csv", unmapped_rows, [
        "source_industry", "occurrences", "example_schemes", "proposed_sector", "review_status",
    ])
    write_csv(reports_out / "unsupported_files.csv", [{"warning": w} for w in warnings], ["warning"])

    quality_summary = {
        "parsed_scheme_records": len(records),
        "records_marked_v1_eligible": sum(1 for r in records if r.get("include_v1")),
        "records_missing_holdings": sum(1 for r in records if "MISSING_HOLDINGS_DATA" in r.get("quality_flags", [])),
        "records_outside_95_101_pct": sum(1 for r in records if "SECTOR_TOTAL_OUTSIDE_95_101" in r.get("quality_flags", [])),
        "records_with_unmapped_equity_industries": sum(1 for r in records if "UNMAPPED_EQUITY_INDUSTRY" in r.get("quality_flags", [])),
        "records_without_scheme_master_match": sum(1 for r in records if r.get("scheme_master_status") == "not_matched"),
        "unmapped_industry_labels": len(unmapped_rows),
        "warnings": warnings,
        "important": "Do not publish funds until the scheme master is populated, mappings are reviewed, and five funds are checked against the original AMC factsheets.",
    }
    (reports_out / "data_quality_summary.json").write_text(json.dumps(quality_summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return quality_summary


def collect_inputs(input_file: Path | None, input_dir: Path | None) -> list[Path]:
    """Find raw Excel workbooks and ZIP archives containing individual AMC workbooks."""
    if input_file:
        if not input_file.exists():
            raise FileNotFoundError(f"Input file not found: {input_file}")
        if input_file.suffix.lower() not in {".xlsx", ".xlsm", ".zip"}:
            raise ValueError("Input must be an .xlsx, .xlsm or .zip file.")
        return [input_file]
    if input_dir is None or not input_dir.exists():
        raise FileNotFoundError(f"Input folder not found: {input_dir}")
    supported = [
        p for p in input_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in {".xlsx", ".xlsm", ".zip"}
        and not p.name.startswith("~$")
    ]
    return sorted(set(supported))


def expand_zip_inputs(files: list[Path], temp_dir: Path) -> tuple[list[Path], list[str]]:
    """Extract only Excel workbook members from ZIP archives into a temporary folder."""
    expanded: list[Path] = []
    warnings: list[str] = []
    used_names: set[str] = set()
    for path in files:
        if path.suffix.lower() != ".zip":
            expanded.append(path)
            used_names.add(path.name.lower())
            continue
        try:
            with zipfile.ZipFile(path) as archive:
                members = [
                    info for info in archive.infolist()
                    if not info.is_dir()
                    and Path(info.filename).suffix.lower() in {".xlsx", ".xlsm"}
                    and not Path(info.filename).name.startswith("~$")
                ]
                if not members:
                    warnings.append(f"No .xlsx/.xlsm workbooks found inside ZIP archive: {path.name}")
                    continue
                extracted_count = 0
                for member in members:
                    basename = Path(member.filename).name
                    if not basename:
                        continue
                    output_name = basename
                    if output_name.lower() in used_names:
                        output_name = f"{path.stem}__{basename}"
                    counter = 2
                    candidate = output_name
                    while candidate.lower() in used_names:
                        stem = Path(output_name).stem
                        suffix = Path(output_name).suffix
                        candidate = f"{stem}_{counter}{suffix}"
                        counter += 1
                    output_path = temp_dir / candidate
                    with archive.open(member, "r") as source, output_path.open("wb") as target:
                        target.write(source.read())
                    expanded.append(output_path)
                    used_names.add(candidate.lower())
                    extracted_count += 1
                print(f"Expanded {path.name}: {extracted_count} Excel workbooks")
        except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
            warnings.append(f"Failed to read ZIP archive {path.name}: {type(exc).__name__}: {exc}")
    return expanded, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description="Create sector exposure JSON and data-quality reports from AMC Excel disclosures.")
    source = parser.add_mutually_exclusive_group()
    source.add_argument("--input", type=Path, help="Process one .xlsx/.xlsm workbook (useful for the first test).")
    source.add_argument("--input-dir", type=Path, help=f"Folder to scan (default: {DEFAULT_RAW_DIR}).")
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING, help="Editable source-industry to standard-sector CSV.")
    parser.add_argument("--scheme-master", type=Path, default=DEFAULT_SCHEME_MASTER, help="Scheme metadata/crosswalk CSV. Leave empty until verified.")
    parser.add_argument("--data-out", type=Path, default=DEFAULT_DATA_OUT, help="Output folder for funds.json.")
    parser.add_argument("--reports-out", type=Path, default=DEFAULT_REPORTS_OUT, help="Output folder for quality reports.")
    args = parser.parse_args()

    input_dir = args.input_dir or (None if args.input else DEFAULT_RAW_DIR)
    files = collect_inputs(args.input, input_dir)
    if not files:
        print(f"No .xlsx/.xlsm workbooks or .zip archives found. Put monthly files in: {input_dir or DEFAULT_RAW_DIR}", file=sys.stderr)
        return 2

    mapping = read_mapping(args.mapping)
    master_rows = read_scheme_master(args.scheme_master)
    all_records: list[dict[str, Any]] = []
    warnings: list[str] = []
    unmapped: dict[str, dict[str, Any]] = {}
    with tempfile.TemporaryDirectory(prefix="fmmf_amc_inputs_") as temporary_name:
        expanded_files, archive_warnings = expand_zip_inputs(files, Path(temporary_name))
        warnings.extend(archive_warnings)
        for path in expanded_files:
            try:
                records, file_warnings = parse_supported_workbook(path, mapping, unmapped)
                if not records and not file_warnings:
                    file_warnings = [f"No scheme records found in {path.name}"]
                all_records.extend(records)
                warnings.extend(file_warnings)
                print(f"Processed {path.name}: {len(records)} scheme tabs")
            except Exception as exc:
                warnings.append(f"Failed to process {path.name}: {type(exc).__name__}: {exc}")
                print(f"WARNING: {warnings[-1]}", file=sys.stderr)

    if not all_records:
        print("No supported AMC portfolios were parsed. See reports/unsupported_files.csv or the warning above.", file=sys.stderr)
        # Still write a warning report to help diagnose format differences.
        write_csv(args.reports_out / "unsupported_files.csv", [{"warning": w} for w in warnings], ["warning"])
        return 1

    summary = build_outputs(all_records, master_rows, unmapped, warnings, args.data_out, args.reports_out)
    print("\nBuild complete")
    print(f"Funds JSON: {args.data_out / 'funds.json'}")
    print(f"Data-quality report: {args.reports_out / 'data_quality.csv'}")
    print(f"Unmapped industries: {args.reports_out / 'unmapped_industries.csv'}")
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
